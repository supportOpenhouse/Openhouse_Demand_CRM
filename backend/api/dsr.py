"""Daily Sales Report (DSR) — one row per PM, grouped area -> city -> total.

Spec: Demand plan July/dsr/DECISIONS.md. Month-to-date up to report date D (IST),
plus a "today" figure for CP meetings, visits and closures.

Sources (each outside DB is optional; a failure only nulls its own columns):
  crm       — this app's DB (visits, followups, property_assignments, users)
  core      — CORE_DATABASE_URL, oh_schedulevisit (assisted %)
  meetings  — MEETINGS_DATABASE_URL, meetings (CP engagement meetings by RM)
  bookings  — PROPERTIES_DATABASE_URL, booking_details (closures)

Settings tables (migration 023) may not exist yet: reads fall back to the built-in
October 2026 defaults below; writes return 409.
"""
from __future__ import annotations

import datetime as dt
import io
import logging
import re
from collections import defaultdict
from typing import Optional

import asyncpg
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from . import auth, config
from .db import acquire

log = logging.getLogger("dsr")
router = APIRouter()

IST = dt.timezone(dt.timedelta(hours=5, minutes=30))
SINCE_AUG = dt.date(2026, 8, 1)

# ---- built-in October 2026 defaults (used until migration 023 runs) ----------
DEFAULT_TARGETS = [  # (city, area, target) — total 100
    ("Gurgaon", "Dwarka Expressway", 10), ("Gurgaon", "New Gurgaon 2", 10),
    ("Gurgaon", "New Gurgaon 1", 10), ("Gurgaon", "SPR", 4),
    ("Noida", "Noida Extension", 23), ("Noida", "Central Noida", 8),
    ("Noida", "Noida Expressway", 7), ("Noida", "Greater Noida", 6),
    ("Ghaziabad", "Raj Nagar Extension", 10), ("Ghaziabad", "Crossings Republik", 6),
    ("Ghaziabad", "NH24", 6),
]
DEFAULT_AREA_MANAGERS = {  # area -> CRM user name (exact, active)
    "Dwarka Expressway": "Ayush Ojha", "New Gurgaon 2": "Ayush Ojha",
    "New Gurgaon 1": "Puran Kiraula", "SPR": "Puran Kiraula",
    "Noida Extension": "Ajitesh Singh",
    "Noida Expressway": "Amar", "Greater Noida": "Amar",
    "Raj Nagar Extension": "Aditya Bhasker",
    "Central Noida": "Manish Pal", "Crossings Republik": "Manish Pal", "NH24": "Manish Pal",
}
AREA_ALIASES = {
    "greater noida west": "Noida Extension",
    "sector 10 greater noida west": "Noida Extension",
    "nh24, ghaziabad": "NH24",
}
CITY_ORDER = ["Gurgaon", "Noida", "Ghaziabad"]
AREA_ORDER = [a for _, a, _ in DEFAULT_TARGETS]
CITY_OF_AREA = {a: c for c, a, _ in DEFAULT_TARGETS}

PCT_FIELDS = {  # pct field -> (numerator key, denominator key)
    "assisted_pct": ("_assisted_n", "_assisted_d"),
    "fu48_pct": ("_fu48", "visits"),
    "fu_any_pct": ("_fu_any", "visits"),
    "revisit_sched_pct": ("_rv_sched", "visits"),
    "revisit_done_pct": ("_rv_done", "visits"),
    "neg_sched_pct": ("_neg_sched", "visits"),
    "neg_held_pct": ("_neg_held", "visits"),
    "dead_pct": ("_dead", "visits"),
    "visit_to_closure_pct": ("closures", "visits"),
    "ach_pct": ("closures", "target"),
}
COUNT_FIELDS = ["cp_meetings", "cp_meetings_today", "visits", "visits_today",
                "closures", "closures_today", "closures_since_aug",
                "_assisted_n", "_assisted_d", "_fu48", "_fu_any", "_rv_sched",
                "_rv_done", "_neg_sched", "_neg_held", "_dead"]


def norm_area(raw: Optional[str]) -> Optional[str]:
    s = (raw or "").strip()
    if not s:
        return None
    return AREA_ALIASES.get(s.lower(), s)


def _dsn(url: str) -> str:
    # asyncpg doesn't accept libpq's channel_binding param (Neon adds it).
    return re.sub(r"[?&]channel_binding=[^&]*", "", url or "")


async def _ext_connect(url: str) -> asyncpg.Connection:
    return await asyncpg.connect(_dsn(url), timeout=15, statement_cache_size=0)


# ---- settings ----------------------------------------------------------------
async def load_settings(conn, month: str, users: dict) -> dict:
    """Targets / area managers / PM managers for `month` ('YYYY-MM').
    Falls back to the October defaults when the tables don't exist yet."""
    by_name = {}
    for u in users.values():
        if u["active"]:
            by_name.setdefault(u["name"].strip().lower(), u["id"])
    targets = {a: {"city": c, "target": t} for c, a, t in DEFAULT_TARGETS}
    area_mgr = {a: by_name.get(n.lower()) for a, n in DEFAULT_AREA_MANAGERS.items()}
    pm_mgr: dict = {}
    log_rows: list = []
    # to_regclass check (not try/except) so a missing table never aborts the caller's txn
    tables = await conn.fetchval(
        "SELECT to_regclass('dsr_area_targets') IS NOT NULL "
        "AND to_regclass('dsr_change_log') IS NOT NULL")
    if tables:
        rows = await conn.fetch(
            "SELECT city, area, target FROM dsr_area_targets WHERE month = $1", month)
        if rows:
            targets = {r["area"]: {"city": r["city"], "target": float(r["target"])
                                   if r["target"] is not None else None} for r in rows}
        rows = await conn.fetch("SELECT area, manager_user_id FROM dsr_area_managers")
        for r in rows:
            area_mgr[r["area"]] = str(r["manager_user_id"]) if r["manager_user_id"] else None
        rows = await conn.fetch("SELECT pm_user_id, manager_user_id FROM dsr_pm_managers")
        pm_mgr = {str(r["pm_user_id"]): (str(r["manager_user_id"]) if r["manager_user_id"] else None)
                  for r in rows}
        rows = await conn.fetch(
            "SELECT l.at, u.name AS by_name, l.kind, l.key, l.old_value, l.new_value "
            "FROM dsr_change_log l LEFT JOIN users u ON u.id = l.by_user_id "
            "ORDER BY l.at DESC LIMIT 200")
        log_rows = [{"at": r["at"].isoformat(), "by": r["by_name"], "kind": r["kind"],
                     "key": r["key"], "old": r["old_value"], "new": r["new_value"]} for r in rows]
    for t in targets.values():  # whole-number targets read back as int
        if isinstance(t["target"], float) and t["target"].is_integer():
            t["target"] = int(t["target"])
    return {"targets": targets, "area_mgr": area_mgr, "pm_mgr": pm_mgr,
            "log": log_rows, "tables": tables}


async def _users(conn) -> dict:
    rows = await conn.fetch(
        "SELECT id, name, team, active, core_sales_manager_id, lower(email) AS email FROM users")
    return {str(r["id"]): {"id": str(r["id"]), "name": r["name"], "team": r["team"],
                           "active": r["active"], "core_id": r["core_sales_manager_id"], "email": r["email"]}
            for r in rows}


# ---- CRM: PM main areas and per-home areas -----------------------------------
async def _pm_areas(conn):
    """property_id -> (area, city, home_id, supply_form_uid); pm -> main area."""
    props = {}
    for r in await conn.fetch(
            "SELECT id, micro_market, city, home_id, supply_form_uid FROM properties"):
        props[str(r["id"])] = {"area": norm_area(r["micro_market"]), "city": r["city"],
                               "home_id": r["home_id"], "uid": r["supply_form_uid"]}
    counts = defaultdict(lambda: defaultdict(int))
    for r in await conn.fetch("SELECT property_id, pm_user_id FROM v_property_current_pm"):
        a = props.get(str(r["property_id"]), {}).get("area")
        if a:
            counts[str(r["pm_user_id"])][a] += 1
    main = {pm: max(c.items(), key=lambda kv: (kv[1], kv[0]))[0] for pm, c in counts.items()}
    area_city = {}
    for p in props.values():
        if p["area"] and p["area"] not in area_city:
            area_city[p["area"]] = CITY_OF_AREA.get(p["area"], p["city"] or "Other")
    area_city.update(CITY_OF_AREA)
    return props, main, area_city


async def _pm_on(conn, pairs: list) -> dict:
    """[(property_id, date)] -> {(property_id, date): pm_user_id} using assignment
    history at noon IST that day; fallback current PM."""
    if not pairs:
        return {}
    pids = list({p for p, _ in pairs})
    hist = defaultdict(list)
    for r in await conn.fetch(
            "SELECT property_id, pm_user_id, effective_from, effective_to "
            "FROM property_assignments WHERE property_id = ANY($1::uuid[])", pids):
        hist[str(r["property_id"])].append(r)
    cur = {}
    for r in await conn.fetch(
            "SELECT property_id, pm_user_id FROM v_property_current_pm "
            "WHERE property_id = ANY($1::uuid[])", pids):
        cur[str(r["property_id"])] = str(r["pm_user_id"])
    out = {}
    day = lambda x: x.astimezone(IST).date() if x is not None else None
    for pid, d in pairs:
        # PM who held the home on day d (IST): latest assignment starting on or before d and not ended before d
        best = None
        for r in hist.get(pid, []):
            f, t = day(r["effective_from"]), day(r["effective_to"])
            if f is not None and f <= d and (t is None or t >= d) and (best is None or f > best[0]):
                best = (f, str(r["pm_user_id"]))
        out[(pid, d)] = best[1] if best else cur.get(pid)
    return out


# ---- CRM visits ----------------------------------------------------------------
VISITS_SQL = """
WITH v AS (
  SELECT v.id, v.buyer_id, v.property_id, v.home_id, v.current_stage, v.current_status,
         v.negotiation_date, coalesce(v.visit_date, v.selected_date) AS vd
    FROM visits v
   WHERE v.status = 'completed'
     AND coalesce(v.visit_date, v.selected_date) BETWEEN $1 AND $2
)
SELECT v.*,
  EXISTS (SELECT 1 FROM followups f WHERE f.visit_id = v.id
           AND f.created_at < (v.vd::timestamp + interval '2 days') AT TIME ZONE 'Asia/Kolkata'
           AND f.created_at < $3) AS fu48,
  EXISTS (SELECT 1 FROM followups f WHERE f.visit_id = v.id AND f.created_at < $3) AS fu_any,
  EXISTS (SELECT 1 FROM followups f WHERE f.visit_id = v.id AND f.created_at < $3
           AND (f.stage IN ('revisit_scheduled','after_revisit_fu') OR f.revisit_date IS NOT NULL))
    AS rv_sched,
  EXISTS (SELECT 1 FROM followups f WHERE f.visit_id = v.id AND f.created_at < $3
           AND f.stage IN ('after_negotiation_fu','booking','ats')) AS neg_moved_on,
  EXISTS (SELECT 1 FROM visits p WHERE p.status = 'completed' AND p.id <> v.id
           AND p.buyer_id = v.buyer_id
           AND (p.property_id = v.property_id OR (p.home_id IS NOT NULL AND p.home_id = v.home_id))
           AND coalesce(p.visit_date, p.selected_date) < v.vd) AS rv_done
  FROM v
"""


async def _crm_visits(conn, start: dt.date, d: dt.date) -> list:
    cutoff = dt.datetime.combine(d + dt.timedelta(days=1), dt.time(0), tzinfo=IST)
    rows = await conn.fetch(VISITS_SQL, start, d, cutoff)
    return [dict(r) for r in rows]


# ---- outside sources -----------------------------------------------------------
async def _core_assisted(start: dt.date, d: dt.date):
    """home_id -> [completed, with_feedback] for core visits MTD. None if unavailable."""
    url = getattr(config, "CORE_DATABASE_URL", "")
    if not url:
        return None
    try:
        c = await _ext_connect(url)
        try:
            rows = await c.fetch(
                "SELECT home_id::text AS home_id, count(*) AS n, "
                "count(demand_sm_feedback_id) AS fb FROM oh_schedulevisit "
                "WHERE status = 'completed' AND selected_date BETWEEN $1 AND $2 "
                "GROUP BY 1", start, d)
        finally:
            await c.close()
        return {r["home_id"]: (r["n"], r["fb"]) for r in rows}
    except Exception as e:  # noqa: BLE001
        log.warning("DSR core source failed: %s", e)
        return None


async def _meetings(start: dt.date, d: dt.date):
    """[(core smid, rm email, date)] engagement meetings MTD. None if unavailable."""
    url = config.MEETINGS_DATABASE_URL
    if not url:
        return None
    try:
        c = await _ext_connect(url)
        try:
            rows = await c.fetch(
                "SELECT u.smid, lower(u.email) AS email, (m.started_at AT TIME ZONE 'Asia/Kolkata')::date AS md "
                "FROM meetings m JOIN users u ON u.id = m.rm_id "
                "WHERE m.meeting_type = 'engagement' "
                "AND coalesce(m.status, '') NOT ILIKE $3 "
                "AND (m.started_at AT TIME ZONE 'Asia/Kolkata')::date BETWEEN $1 AND $2",
                start, d, "fail%")
        finally:
            await c.close()
        return [(r["smid"], r["email"], r["md"]) for r in rows]
    except Exception as e:  # noqa: BLE001
        log.warning("DSR meetings source failed: %s", e)
        return None


async def _bookings(d: dt.date):
    """[(uid, booking date)] active bookings from 1 Aug to D. None if unavailable."""
    url = config.PROPERTIES_DATABASE_URL
    if not url:
        return None
    try:
        c = await _ext_connect(url)
        try:
            rows = await c.fetch(
                """
                SELECT DISTINCT ON (b.uid) b.uid,
                       coalesce(b.booking_date::date,
                                (b.created_at AT TIME ZONE 'Asia/Kolkata')::date) AS bd
                  FROM booking_details b
                  JOIN properties p ON p.uid = b.uid
                  JOIN demand_details dd ON dd.uid = b.uid
                 WHERE dd.availability_status IN ('Booked', 'Sold')
                   AND coalesce(b.buyer_name, '') NOT ILIKE $1
                   AND coalesce(b.selling_cp_name, '') NOT ILIKE $1
                   AND coalesce(p.society_name, '') <> 'DLF The Camellias'
                 ORDER BY b.uid, b.created_at
                """, "%test%")
        finally:
            await c.close()
        return [(r["uid"], r["bd"]) for r in rows if r["bd"] and SINCE_AUG <= r["bd"] <= d]
    except Exception as e:  # noqa: BLE001
        log.warning("DSR bookings source failed: %s", e)
        return None


# ---- assembly ------------------------------------------------------------------
def _blank(kind, city, area, pm_id=None, pm_name=None, manager=None, target=None) -> dict:
    r = {"kind": kind, "city": city, "area": area, "pm_id": pm_id, "pm_name": pm_name,
         "manager_name": manager, "target": target}
    for k in COUNT_FIELDS:
        r[k] = 0
    return r


def _finish(r: dict, days: int, sources: dict) -> dict:
    out = {k: v for k, v in r.items() if not k.startswith("_")}
    for f, (n, dkey) in PCT_FIELDS.items():
        den = r.get(dkey)
        out[f] = (r[n] / den) if den else None
    if r["kind"] == "pm":
        out["target"] = None
        out["ach_pct"] = None
    out["weekly_visit_avg"] = round(r["visits"] / (days / 7), 2) if days else None
    if not sources["core"]:
        out["assisted_pct"] = None
    if not sources["meetings"]:
        out["cp_meetings"] = out["cp_meetings_today"] = None
    if not sources["bookings"]:
        for k in ("closures", "closures_today", "closures_since_aug",
                  "visit_to_closure_pct", "ach_pct"):
            out[k] = None
    return out


async def compute_dsr(conn, d: dt.date, user: Optional[dict] = None) -> dict:
    start = d.replace(day=1)
    days = (d - start).days + 1
    month = d.strftime("%Y-%m")
    users = await _users(conn)
    st = await load_settings(conn, month, users)
    props, main_area, area_city = await _pm_areas(conn)

    acc: dict = {}  # pm_id -> counters

    def bucket(pm):
        if pm not in acc:
            acc[pm] = defaultdict(int)
        return acc[pm]

    # visits
    visits = await _crm_visits(conn, start, d)
    # a visit with no property link still names its home (home_id): find the CRM property for that home
    home_to_pid = {str(p["home_id"]): pid for pid, p in props.items() if p["home_id"]}
    for v in visits:
        v["pid"] = str(v["property_id"]) if v["property_id"] else home_to_pid.get(str(v["home_id"]))
    pm_at = await _pm_on(conn, [(v["pid"], v["vd"]) for v in visits if v["pid"]])
    for v in visits:
        pm = pm_at.get((v["pid"], v["vd"])) if v["pid"] else None
        if not pm:
            continue
        b = bucket(pm)
        b["visits"] += 1
        b["visits_today"] += v["vd"] == d
        b["_fu48"] += v["fu48"]
        b["_fu_any"] += v["fu_any"]
        b["_rv_sched"] += v["rv_sched"]
        b["_rv_done"] += v["rv_done"]
        nd = v["negotiation_date"]
        nd = nd.astimezone(IST).date() if nd else None
        neg = v["current_stage"] == "negotiation"
        b["_neg_sched"] += bool(neg and nd and nd >= d)
        b["_neg_held"] += bool(v["neg_moved_on"] or (neg and nd and nd < d)
                               or v["current_stage"] in ("booking", "ats"))
        b["_dead"] += v["current_status"] == "dead"

    # core assisted (by home_id -> the home's PM on the report date)
    core = await _core_assisted(start, d)
    if core is not None:
        hid_to_pid = {str(p["home_id"]): pid for pid, p in props.items() if p["home_id"]}
        cur_pm = await _pm_on(conn, [(hid_to_pid[h], d) for h in core if h in hid_to_pid])
        for h, (n, fb) in core.items():
            pid = hid_to_pid.get(h)
            pm = cur_pm.get((pid, d)) if pid else None
            if pm:
                b = bucket(pm)
                b["_assisted_d"] += n
                b["_assisted_n"] += fb

    # CP meetings -> CRM user by email, else core_sales_manager_id (prefer the active user)
    meets = await _meetings(start, d)
    if meets is not None:
        by_core, by_email = {}, {}
        for u in sorted(users.values(), key=lambda u: not u["active"]):
            if u["core_id"] is not None:
                by_core.setdefault(str(u["core_id"]), u["id"])
            if u["email"]:
                by_email.setdefault(u["email"], u["id"])
        for smid, email, md in meets:
            uid = by_email.get(email) or by_core.get(str(smid))
            if uid:
                b = bucket(uid)
                b["cp_meetings"] += 1
                b["cp_meetings_today"] += md == d

    # closures -> home's PM on booking date (booking uid -> properties.supply_form_uid)
    books = await _bookings(d)
    if books is not None:
        uid_to_pid = {str(p["uid"]): pid for pid, p in props.items() if p["uid"]}
        pairs = [(uid_to_pid[u], bd) for u, bd in books if u in uid_to_pid]
        bpm = await _pm_on(conn, pairs)
        for u, bd in books:
            pid = uid_to_pid.get(u)
            pm = bpm.get((pid, bd)) if pid else None
            if not pm:
                continue
            b = bucket(pm)
            b["closures_since_aug"] += 1
            if bd >= start:
                b["closures"] += 1
                b["closures_today"] += bd == d

    sources = {"crm": True, "core": core is not None, "meetings": meets is not None,
               "bookings": books is not None}

    # group PMs under main area (PMs with no homes: fall back to "Unassigned")
    def mgr_name(uid):
        return users[uid]["name"] if uid and uid in users else None

    by_area = defaultdict(list)
    for pm, b in acc.items():
        if pm not in users:
            continue
        area = main_area.get(pm) or "Unassigned"
        mgr = st["pm_mgr"].get(pm) or st["area_mgr"].get(area)
        row = _blank("pm", area_city.get(area, "Other"), area, pm, users[pm]["name"],
                     mgr_name(mgr))
        for k, v in b.items():
            row[k] = v
        by_area[area].append(row)
    for area in st["targets"]:
        by_area.setdefault(area, [])

    def city_of(area):
        return st["targets"].get(area, {}).get("city") or area_city.get(area, "Other")

    cities = defaultdict(list)
    for area in by_area:
        cities[city_of(area)].append(area)

    def ckey(c):
        return (CITY_ORDER.index(c) if c in CITY_ORDER else 99, c)

    def akey(a):
        return (AREA_ORDER.index(a) if a in AREA_ORDER else 99, a)

    rows = []
    total = _blank("total", None, None, target=0)
    for city in sorted(cities, key=ckey):
        crow = _blank("city", city, None, target=0)
        for area in sorted(cities[city], key=akey):
            tgt = st["targets"].get(area, {}).get("target")
            arow = _blank("area", city, area, manager=mgr_name(st["area_mgr"].get(area)),
                          target=tgt)
            for pr in sorted(by_area[area], key=lambda r: r["pm_name"] or ""):
                pr["city"] = city
                rows.append(pr)
                for k in COUNT_FIELDS:
                    arow[k] += pr[k]
            rows.append(arow)
            for k in COUNT_FIELDS:
                crow[k] += arow[k]
            crow["target"] += tgt or 0
        rows.append(crow)
        for k in COUNT_FIELDS:
            total[k] += crow[k]
        total["target"] += crow["target"]
    rows.append(total)

    is_admin = bool(user and user.get("team") == "Admin")
    uid = str(user["id"]) if user and user.get("id") else None
    can_edit = list(st["targets"]) if is_admin else \
        [a for a, m in st["area_mgr"].items() if uid and m == uid]
    return {"date": d.isoformat(), "month_start": start.isoformat(), "days_elapsed": days,
            "sources": sources, "can_edit_targets": can_edit, "is_admin": is_admin,
            "rows": [_finish(r, days, sources) for r in rows]}


# ---- access ----------------------------------------------------------------------
async def _is_manager(conn, user: dict) -> bool:
    users = await _users(conn)
    st = await load_settings(conn, dt.date.today().strftime("%Y-%m"), users)
    uid = str(user.get("id"))
    return uid in {m for m in st["area_mgr"].values() if m} or \
        uid in {m for m in st["pm_mgr"].values() if m}


async def _require_view(conn, user: dict) -> None:
    if user.get("team") == "Admin":
        return
    if not await _is_manager(conn, user):
        raise HTTPException(403, "DSR is for Admin and area managers")


def _parse_date(s: Optional[str]) -> dt.date:
    if not s:
        return dt.datetime.now(IST).date()
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        raise HTTPException(400, "date must be YYYY-MM-DD")


# ---- endpoints -------------------------------------------------------------------
@router.get("/api/dsr")
async def get_dsr(date: Optional[str] = None, user: dict = Depends(auth.current_user)):
    d = _parse_date(date)
    async with acquire() as conn:
        await _require_view(conn, user)
        return await compute_dsr(conn, d, user)


@router.get("/api/dsr/settings")
async def get_settings(month: Optional[str] = None, user: dict = Depends(auth.current_user)):
    month = month or dt.datetime.now(IST).strftime("%Y-%m")
    async with acquire() as conn:
        await _require_view(conn, user)
        users = await _users(conn)
        st = await load_settings(conn, month, users)
        _, main_area, _ = await _pm_areas(conn)

    def nm(u):
        return users[u]["name"] if u and u in users else None
    pms = sorted({pm for pm in main_area if pm in users} | set(st["pm_mgr"]),
                 key=lambda u: nm(u) or "")
    return {
        "month": month,
        "area_targets": [{"city": t["city"], "area": a, "target": t["target"]}
                         for a, t in st["targets"].items()],
        "area_managers": [{"area": a, "manager_user_id": m, "manager_name": nm(m)}
                          for a, m in st["area_mgr"].items()],
        "pm_managers": [{"pm_user_id": p, "pm_name": nm(p),
                         "manager_user_id": st["pm_mgr"].get(p),
                         "manager_name": nm(st["pm_mgr"].get(p))} for p in pms],
        "users": [{"id": u["id"], "name": u["name"], "team": u["team"]}
                  for u in sorted(users.values(), key=lambda u: u["name"]) if u["active"]],
        "log": st["log"],
    }


class TargetBody(BaseModel):
    month: str
    area: str
    target: Optional[float] = None


class AreaManagerBody(BaseModel):
    area: str
    manager_user_id: Optional[str] = None


class PmManagerBody(BaseModel):
    pm_user_id: str
    manager_user_id: Optional[str] = None


async def _write(conn, user, kind, key, old, new, sql, *args):
    try:
        async with conn.transaction():
            await conn.execute(sql, *args)
            await conn.execute(
                "INSERT INTO dsr_change_log (by_user_id, kind, key, old_value, new_value) "
                "VALUES ($1, $2, $3, $4, $5)", user["id"], kind, key,
                None if old is None else str(old), None if new is None else str(new))
    except asyncpg.UndefinedTableError:
        raise HTTPException(409, "settings tables not created yet")


@router.put("/api/dsr/target")
async def put_target(body: TargetBody, user: dict = Depends(auth.current_user)):
    async with acquire() as conn:
        users = await _users(conn)
        st = await load_settings(conn, body.month, users)
        if not st["tables"]:
            raise HTTPException(409, "settings tables not created yet")
        if user.get("team") != "Admin" and st["area_mgr"].get(body.area) != str(user["id"]):
            raise HTTPException(403, "Only Admin or this area's manager can edit its target")
        old = st["targets"].get(body.area, {})
        city = old.get("city") or CITY_OF_AREA.get(body.area)
        await _write(conn, user, "target", f"{body.month}|{body.area}", old.get("target"),
                     body.target,
                     "INSERT INTO dsr_area_targets (month, city, area, target, updated_by, updated_at) "
                     "VALUES ($1, $2, $3, $4, $5, now()) ON CONFLICT (month, area) DO UPDATE SET "
                     "target = EXCLUDED.target, updated_by = EXCLUDED.updated_by, updated_at = now()",
                     body.month, city, body.area, body.target, user["id"])
    return {"ok": True}


async def _require_admin_or_manager(conn, user) -> None:
    if user.get("team") != "Admin" and not await _is_manager(conn, user):
        raise HTTPException(403, "Admin or area managers only")


@router.put("/api/dsr/area-manager")
async def put_area_manager(body: AreaManagerBody, user: dict = Depends(auth.current_user)):
    async with acquire() as conn:
        await _require_admin_or_manager(conn, user)
        users = await _users(conn)
        st = await load_settings(conn, dt.datetime.now(IST).strftime("%Y-%m"), users)
        if not st["tables"]:
            raise HTTPException(409, "settings tables not created yet")
        await _write(conn, user, "area_manager", body.area, st["area_mgr"].get(body.area),
                     body.manager_user_id,
                     "INSERT INTO dsr_area_managers (area, manager_user_id, updated_by, updated_at) "
                     "VALUES ($1, $2::uuid, $3, now()) ON CONFLICT (area) DO UPDATE SET "
                     "manager_user_id = EXCLUDED.manager_user_id, updated_by = EXCLUDED.updated_by, "
                     "updated_at = now()", body.area, body.manager_user_id, user["id"])
    return {"ok": True}


@router.put("/api/dsr/pm-manager")
async def put_pm_manager(body: PmManagerBody, user: dict = Depends(auth.current_user)):
    async with acquire() as conn:
        await _require_admin_or_manager(conn, user)
        users = await _users(conn)
        st = await load_settings(conn, dt.datetime.now(IST).strftime("%Y-%m"), users)
        if not st["tables"]:
            raise HTTPException(409, "settings tables not created yet")
        await _write(conn, user, "pm_manager", body.pm_user_id, st["pm_mgr"].get(body.pm_user_id),
                     body.manager_user_id,
                     "INSERT INTO dsr_pm_managers (pm_user_id, manager_user_id, updated_by, updated_at) "
                     "VALUES ($1::uuid, $2::uuid, $3, now()) ON CONFLICT (pm_user_id) DO UPDATE SET "
                     "manager_user_id = EXCLUDED.manager_user_id, updated_by = EXCLUDED.updated_by, "
                     "updated_at = now()", body.pm_user_id, body.manager_user_id, user["id"])
    return {"ok": True}


# ---- Excel export ------------------------------------------------------------------
XLSX_COLS = [  # (group, header, field, kind)
    ("IDENTITY", "City", "city", "t"), ("IDENTITY", "Area", "area", "t"),
    ("IDENTITY", "PM", "pm_name", "t"), ("IDENTITY", "Manager", "manager_name", "t"),
    ("TARGETS", "Target", "target", "n"), ("TARGETS", "Ach %", "ach_pct", "p"),
    ("ACTIVITY", "CP Meetings", "cp_meetings", "n"), ("ACTIVITY", "CP Meetings Today", "cp_meetings_today", "n"),
    ("ACTIVITY", "Weekly Visit Avg", "weekly_visit_avg", "d"), ("ACTIVITY", "Visits", "visits", "n"),
    ("ACTIVITY", "Visits Today", "visits_today", "n"),
    ("QUALITY", "Assisted %", "assisted_pct", "p"), ("QUALITY", "FU <48h %", "fu48_pct", "p"),
    ("QUALITY", "Any FU %", "fu_any_pct", "p"), ("QUALITY", "Revisit Sched %", "revisit_sched_pct", "p"),
    ("QUALITY", "Revisit Done %", "revisit_done_pct", "p"), ("QUALITY", "Neg Sched %", "neg_sched_pct", "p"),
    ("QUALITY", "Neg Held %", "neg_held_pct", "p"),
    ("OUTCOME", "Closures", "closures", "n"), ("OUTCOME", "Closures Today", "closures_today", "n"),
    ("OUTCOME", "Visit to Closure %", "visit_to_closure_pct", "p"),
    ("EFFICIENCY", "Visit to Dead %", "dead_pct", "p"),
    ("EFFICIENCY", "Closures since 1 Aug", "closures_since_aug", "n"),
]


def build_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "DSR " + report["date"]
    bold = Font(bold=True)
    # row 1: group headers, merged across each group's columns
    col = 1
    groups = []
    for g, _, _, _ in XLSX_COLS:
        if groups and groups[-1][0] == g:
            groups[-1][2] = col
        else:
            groups.append([g, col, col])
        col += 1
    for g, c1, c2 in groups:
        ws.cell(row=1, column=c1, value=g).font = bold
        ws.cell(row=1, column=c1).alignment = Alignment(horizontal="center")
        if c2 > c1:
            ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
    for i, (_, h, _, _) in enumerate(XLSX_COLS, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = bold
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    fill = PatternFill("solid", fgColor="EEEEEE")
    r = 3
    for row in report["rows"]:
        label = {"area": "Total", "city": "Total", "total": "Grand Total"}.get(row["kind"])
        for i, (_, _, f, k) in enumerate(XLSX_COLS, 1):
            v = row.get(f)
            if f == "pm_name" and label:
                v = label
            c = ws.cell(row=r, column=i, value=v)
            if k == "p":
                c.number_format = "0%"
            elif k == "d":
                c.number_format = "0.0"
            if row["kind"] != "pm":
                c.font = bold
                c.fill = fill
        r += 1
    for i in range(1, len(XLSX_COLS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18 if i <= 4 else 11
    ws.freeze_panes = "E3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/api/dsr/export.xlsx")
async def export_xlsx(date: Optional[str] = None, user: dict = Depends(auth.current_user)):
    d = _parse_date(date)
    async with acquire() as conn:
        await _require_view(conn, user)
        report = await compute_dsr(conn, d, user)
    data = build_xlsx(report)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="DSR_{d.isoformat()}.xlsx"'})
